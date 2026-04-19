from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = Document()
    
    # Title
    title = doc.add_heading('UNSW-NB15 IDS Boru Hattı Optimizasyon Raporu', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Intro
    doc.add_paragraph(
        "Bu rapor, UNSW-NB15 veri seti üzerinde %65 Macro F1-skoru başarısına ulaşan "
        "saldırı tespit sisteminin (IDS) uçtan uca mimarisini, kullanılan siber güvenlik "
        "odaklı metotları ve karşılaştırmalı sonuçları teknik ayrıntılarıyla sunar."
    )

    # Section 1
    doc.add_heading('1. Veri Edinme ve Gelişmiş Önbellekleme', level=1)
    p1 = doc.add_paragraph()
    p1.add_run('Veri Yükleme: ').bold = True
    p1.add_run('UNSW-NB15 ana CSV dosyaları, NUSW-NB15_features.csv meta-verisi ile eşleştirilerek yüklenir.\n')
    p1.add_run('Hibrit Özellik Seti: ').bold = True
    p1.add_run('Modeller hem orijinal (36) hem de geliştirilmiş (55) özellik setleri üzerinde paralel olarak eğitilir.\n')
    p1.add_run('Performans Önbelleği: ').bold = True
    p1.add_run('İşlenen veri yapıları joblib ile disk üzerinde (Xt_tr.joblib, vb.) saklanır.')

    # Section 2
    doc.add_heading('2. Siber Güvenlik Odaklı Cerrahi Özellik Mühendisliği', level=1)
    doc.add_paragraph(
        "Saldırganların davranışsal izlerini (TTP) yakalamak için orijinal 36 özelliğe ek olarak "
        "19 yeni cerrahi özellik geliştirilmiştir (Toplam 55 kolon)."
    )
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Kategori'
    hdr_cells[1].text = 'Özellik Sayısı'
    hdr_cells[2].text = 'Teknik Gerekçe'
    
    data = [
        ['Port Analizi', '7', 'Arka kapıların (Backdoor) yüksek port kullanımı (>49152) ve HTTP/DNS taklidi tespiti.'],
        ['Trafik Yoğunluğu', '4', 'DoS saldırılarındaki yüksek frekansta küçük paket gönderimi tespiti.'],
        ['Trafik Asimetrisi', '3', 'Tarama ve Analiz faaliyetlerindeki asimetrik veri akışı tespiti.'],
        ['TTL & Zamanlama', '5', 'TTL anomali analizi ve Jitter oranları ile Bot/Script imzaları.'],
    ]
    
    for cat, count, reason in data:
        row_cells = table.add_row().cells
        row_cells[0].text = cat
        row_cells[1].text = count
        row_cells[2].text = reason

    # Section 3
    doc.add_heading('3. Hibrit Veri Yeniden Örnekleme Strategy', level=1)
    doc.add_paragraph("1. Akıllı Alt-Örnekleme: Normal sınıfı 200k örneğe düşürüldü.")
    doc.add_paragraph("2. Sınır Temizliği (Tomek Links): Gürültülü sınırlar temizlendi.")
    doc.add_paragraph("3. Sentetik Üst-Örnekleme (SMOTE): Nadir saldırılar (Worms, vb.) artırıldı.")

    # Section 4/5/6/7 combined summary for brevity or separate sections
    doc.add_heading('4. Model Bazlı Karşılaştırmalı Sonuçlar (Test Seti)', level=1)
    
    res_table = doc.add_table(rows=1, cols=6)
    res_table.style = 'Table Grid'
    hdr = res_table.rows[0].cells
    hdr[0].text = 'Sınıf'
    hdr[1].text = 'XGBoost'
    hdr[2].text = 'LGBM (Ori)'
    hdr[3].text = 'LGBM V2'
    hdr[4].text = 'HistGB'
    hdr[5].text = 'Ensemble'
    
    results = [
        ['Analysis', '0.1647', '0.1822', '0.1515', '0.1382', '0.1756'],
        ['Backdoor', '0.1144', '0.1186', '0.0680', '0.1120', '0.1155'],
        ['DoS', '0.3400', '0.3351', '0.2573', '0.4402', '0.3919'],
        ['Worms', '0.5714', '0.6190', '0.6809', '0.5176', '0.7347'],
        ['Macro F1', '0.6259', '0.6345', '0.6215', '0.6082', '0.6473'],
    ]
    
    for row in results:
        cells = res_table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val

    doc.add_heading('5. Akıllı Topluluk (Heuristic Ensemble v5)', level=1)
    doc.add_paragraph(
        "Sistemin nihai kararı, basit bir ortalama yerine 'Güvenilir Uzman Yönlendirmesi' "
        "ile verilir. Bu hibrit yapı, bireysel modellerin en güçlü olduğu saldırı türlerine "
        "göre tahmini dinamik olarak yönlendirir. Sonuç: %64.73 Macro F1 başarısı."
    )

    doc.save('ids_pipeline_report.docx')
    print("ids_pipeline_report.docx başarıyla oluşturuldu.")

if __name__ == "__main__":
    create_report()
